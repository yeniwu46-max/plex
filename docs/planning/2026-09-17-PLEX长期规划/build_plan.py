from pathlib import Path
from datetime import datetime
import math
import unicodedata
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

OUT = Path(__file__).parent
ROOT = Path('D:/zhongruan/zhongruan')
FILE = OUT / 'PLEX_长期完善与落地计划_5表.xlsx'
wb = Workbook()
wb.remove(wb.active)
NAVY, TEAL, LIGHT, INK, GRAY = '17324D', '087F8C', 'EAF4F5', '23364B', '687889'
YELLOW, BLUE = 'FFF2CC', '0000FF'
DATE = 'yyyy-mm-dd'
NUM = '#,##0;(#,##0);"—"'
MONEY = '¥#,##0.00;(¥#,##0.00);"—"'
PCT = '0.0%;(0.0%);"—"'
SOURCES = {
 'U': 'Source: 用户说明与确认，2026-09-17：公司、4班约150人测试、模型微调、1000元营收均已完成；本文件未审计原始凭证。',
 'K': f'Source: 仓库课程知识库，读取于2026-09-17，{ROOT}/backend/data/rag_docs/python-basics/README.md；文档记载8模块、100知识点，不等于已完成教学质量验证。',
 'P': f'Source: 仓库个性化实现说明（2026-06-08），{ROOT}/docs/2026-06-08-personalization-resource-implementation.md；只采纳流程描述，历史16知识点计数已被后续文档更新。',
 'H': f'Source: 硬件状态文档（截至2026-08-31），{ROOT}/hardware/plex-edge/STATUS.md；最新实物状态尚待项目组更新。',
 'B': f'Source: 内部BOM估算，{ROOT}/hardware/plex-edge/docs/BOM.csv，读取于2026-09-17，234元/套；非供应商报价，未含人工、开模、检测和税费。',
 'PY': 'Source: Python官方教程，访问2026-09-17，https://docs.python.org/zh-cn/3/tutorial/ ，用于课程覆盖校对；教学顺序为本计划建议。',
 'WWC': 'Source: What Works Clearinghouse Handbook v5.0，访问2026-09-17，https://ies.ed.gov/ncee/wwc/handbooks ，用于基线、流失与班级聚类的研究设计参考。',
 'AL': 'Source: Freeman et al., 2014, PNAS，https://doi.org/10.1073/pnas.1319030111 ，主动学习设计依据；不是PLEX效果的实测证据。',
 'MC': 'Source: Mitchell et al., Model Cards for Model Reporting，https://research.google/pubs/model-cards-for-model-reporting/ ，访问2026-09-17，用于模型用途、性能和局限记录。',
 'LAW': 'Source: 个人信息保护法（官方全文），https://www.cac.gov.cn/2021-08/20/c_1631050028355286.htm ，访问2026-09-17，用于目的限定、最小必要与个人信息权益设计。',
 'PAY': 'Source: 市场监管总局《网络交易监督管理办法》2025年页面，https://www.samr.gov.cn/zw/zfxxgk/fdzdgknr/fgs/art/2025/art_4b47c79b8d994a42bba4835997688faa.html ，访问2026-09-17，用于收费披露与续费服务设计。',
 'PLAN': 'Source: 本计划建议，2026-09-17。目标、阈值、工期、角色和价格均为规划假设，不代表已有结果、真实报价或已签订单。'
}

def comment(cell, *keys):
    cell.comment = Comment('\n\n'.join(SOURCES[k] for k in keys), 'PLEX 规划依据')

def sheet(name, title, subtitle, widths):
    ws = wb.create_sheet(name)
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    n = len(widths)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    ws.cell(1, 1, title).font = Font(name='微软雅黑', size=22, bold=True, color='FFFFFF')
    ws.cell(1, 1).fill = PatternFill('solid', fgColor=NAVY)
    ws.cell(1, 1).alignment = Alignment(vertical='center', indent=1)
    ws.row_dimensions[1].height = 46
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
    ws.cell(2, 1, subtitle).font = Font(name='微软雅黑', size=10, color=GRAY)
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[2].height = 36
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 75
    ws.freeze_panes = 'C6'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = TEAL
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = .25
    ws.page_margins.right = .25
    ws.page_margins.top = .35
    ws.page_margins.bottom = .35
    ws.print_title_rows = '1:2'
    ws.oddFooter.center.text = 'PLEX · 12个月执行 / 24个月展望 · 规划目标≠已实现结果'
    ws.oddFooter.right.text = '&P / &N'
    return ws

def band(ws, row, text, n=None):
    n = n or ws.max_column
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n)
    c = ws.cell(row, 1, text)
    c.fill = PatternFill('solid', fgColor=TEAL)
    c.font = Font(name='微软雅黑', size=12, bold=True, color='FFFFFF')
    c.alignment = Alignment(vertical='center', indent=1)
    ws.row_dimensions[row].height = 28

def header(ws, row, vals):
    for i, val in enumerate(vals, 1):
        c = ws.cell(row, i, val)
        c.fill = PatternFill('solid', fgColor=NAVY)
        c.font = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[row].height = 29

def textwidth(s):
    return sum(2 if unicodedata.east_asian_width(c) in 'WF' else 1 for c in str(s))

def rows(ws, start, data, minheight=46):
    for r, vals in enumerate(data, start):
        maxlines = 1
        for col, val in enumerate(vals, 1):
            c = ws.cell(r, col, val)
            c.font = Font(name='微软雅黑', size=10, color=INK)
            c.fill = PatternFill('solid', fgColor='F2F6FA' if (r-start)%2 == 0 else 'FFFFFF')
            c.alignment = Alignment(wrap_text=True, vertical='center')
            c.border = Border(bottom=Side(style='hair', color='DCE4ED'))
            width = ws.column_dimensions[get_column_letter(col)].width - 2
            if not isinstance(val, str) or not val.startswith('='):
                lines = sum(max(1, math.ceil(textwidth(line)/width)) for line in str(val or '').split('\n'))
                maxlines = max(maxlines, lines)
            if isinstance(val, str) and val.startswith('='):
                c.font = Font(name='微软雅黑', size=10, color='008000' if '!' in val else '000000')
        ws.row_dimensions[r].height = max(minheight, maxlines*15+10)

def note(ws, row, text, n=None, fill='FFF2CC'):
    n = n or ws.max_column
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n)
    c = ws.cell(row, 1, text)
    c.font = Font(name='微软雅黑', size=10, color=INK)
    c.fill = PatternFill('solid', fgColor=fill)
    c.alignment = Alignment(wrap_text=True, vertical='center', indent=1)
    ws.row_dimensions[row].height = 34

# 01: targets are an editable proposal, baseline is the user's latest account.
ws = sheet('01目标路线', 'PLEX｜长期完善与落地计划', '总目标：把已有试用、微调和营收，沉淀为可教、可解释、可复核、可运营的 Python 数字教材产品。12个月执行，24个月条件式扩展。', [17,22,32,34,21,23,24,32])
rows(ws, 4, [['启动日',datetime(2026,9,17),'12个月截止','=EDATE(B4,12)-1','任务数',"=COUNTA('05执行证据'!A6:A37)",'任务完成率',"=IFERROR(COUNTIF('05执行证据'!I6:I37,\"已完成\")/F4,0)"]],32)
ws['B4'].number_format = ws['D4'].number_format = DATE
ws['B4'].font = Font(name='微软雅黑', size=10, color=BLUE)
ws['B4'].fill = PatternFill('solid', fgColor=YELLOW)
ws['H4'].number_format = PCT
comment(ws['B4'],'PLAN')
band(ws,5,'A｜当前基础：按用户最新确认；旧文档只用于定位已有设计')
header(ws,6,['项目','当前值 / 状态','口径','下一步补齐','主责角色','时间','依据','判断边界'])
baseline = [
 ['真实测试班级',4,'已完成4个班测试','班级、课程、时间与教师证明','教研负责人','M1','用户2026-09-17确认','不把4个班当150个独立实验单元'],
 ['真实测试学生',150,'约150人；待按匿名ID去重','参与、有效前后测、退出人数分开列','科研负责人','M1','用户2026-09-17确认','不与人次、事件数混用'],
 ['已有营收（元）',1000,'币种暂按人民币；收入已实现','订单、收款、退款、履约与会计口径对账','财务运营','M1','用户2026-09-17确认','不是利润；待核实净退款口径'],
 ['公司主体',1,'已成立公司','主体、合同、账户、知识产权归属','项目负责人','M1','用户2026-09-17确认','不安排重复注册公司'],
 ['模型微调','已完成','处理数据并完成训练微调','数据版本、模型版本、基线与测试报告','算法负责人','M1—M3','用户2026-09-17确认','性能增益数值待补，不预设优于基座'],
 ['教学与硬件','已有课程库 / 硬件方案','仓库记载8模块100知识点；硬件为设计版','核对教材质量与最新实物进展','教研 / 硬件','M1','课程README / 硬件STATUS','旧文档不是当前实物验收结论'],
 ['收费与公益','已有营收 / 半价活动方向','延续现有收费与学生优惠安排','支付状态、半价规则、受益台账与成本','运营负责人','M1—M2','用户原始描述','活动的实际人数、次数和成本待补']
]
rows(ws,7,baseline)
for r in range(7,12): comment(ws.cell(r,2),'U')
comment(ws['B12'],'K','H')
comment(ws['B13'],'U')
band(ws,15,'B｜12个月验收目标（建议值；以质量与可复制为先）')
header(ws,16,['维度','关键成果','交付与数量','验收口径','完成点','主责角色','任务关联','未达标时的动作'])
goals = [
 ['教育教学','Python教材V1→V2','8单元 / 建议32课时；2个综合项目','目标—活动—评价全映射；100知识点双人审稿','M3 / M9','教研负责人','T02—T04、T30','删减未审内容，优先可完整授课'],
 ['个性化科学性','证据驱动的生成闭环','诊断→检索→生成→审核→学习→更新','抽样100条资源；依据可追溯100%，教师合格率≥90%','M3','产品 / 算法','T05—T06','降低自动发布范围，回退到审定资源'],
 ['科研可信度','可复核的研究结论','现有试验回顾1份；新队列验证1轮；报告1份','预先确定主指标；效应量+95%区间+缺失/局限','M7—M8','科研负责人','T07—T13','如实报告无效/负向结果，不承诺显著'],
 ['模型技术','可复现且受控的模型版本','已有微调复现；固定独立测试集1套','同条件比较基座、RAG、RAG+微调；质量成本一起看','M3—M6','算法负责人','T08—T09、T11','无稳定收益则沿用更便宜的基线'],
 ['软硬件结合','课堂闭环与教学数据实验','1套可用样机→最多10套共享终端','72小时连续运行；事件成功率≥99%；比较软硬件增益','M3—M6','硬件 / 后端','T14—T18','无教学或交互收益则保留软件方案'],
 ['可落地交付','教师可独立使用','新环境部署包、教师手册与恢复演练各1套','连续4周可用率≥99%；教师独立完成1轮课堂','M3—M5','工程负责人','T19—T20、T29','先修阻断项，再招新班'],
 ['商业持续性','从首笔验证到可重复付费','12个月累计净收款暂定≥1万元；2个机构付费单','先统一历史1000元口径；连续3个月经营贡献为正','M9—M12','财务 / 运营','T21—T24、T27、T31','按表04调整价格与成本；不等同净利润'],
 ['规模与公益','受益可量化、扩展可控制','累计≥8班 / ≥300名去重学生；新增≥40半价+20免费名额','新增付费与测试分别统计；免费/半价受益不重复计数','M12','教研 / 运营','T25—T28','先核算补贴承受力，再扩大活动']
]
rows(ws,17,goals)
for r in range(17,25): comment(ws.cell(r,3),'PLAN')
band(ws,26,'C｜阶段关口：任何规模扩展都以前一阶段验收为条件')
header(ws,27,['阶段','建议窗口','必须做成','放行条件','不做什么','负责人','对应任务','课程交付定位'])
roadmap = [
 ['G1 基线固化','M1｜9/17—10/16','归档已有成果；明确课程、数据与收费口径','4类已有成果各有负责人和证据入口','不重复开展已完成事项','项目负责人','T01、T07、T14、T20、T22','先形成可信项目说明'],
 ['G2 可教学交付','M2—M3｜至12/16','教材V1、生成链、模型卡、样机与收费审计','一节完整课可演示；核心流程无阻断','不以界面截图替代课堂闭环','教研 / 工程','T02—T06、T08—T09、T15—T16、T19、T21、T23、T25、T30','课程答辩最小完整包；DDL变化只调整此关口'],
 ['G3 实证与运营','M4—M6｜至2027/3/16','新队列验证、硬件试点、真实付费与公益台账','完成预定干预及数据冻结；订单可逐笔对账','不把四班结果外推至所有学生','科研 / 运营','T10—T11、T17—T18、T24、T26、T29','补充真实课堂案例'],
 ['G4 复制与复盘','M7—M12｜至2027/9/16','报告、第二机构订单、教材V2与年度复盘','累计8班300人；正经营贡献；证据包可复核','不以亏损补贴长期换人数','项目负责人','T12—T13、T27—T28、T31—T32','可复制的课程产品'],
 ['24个月方向','M13—M24｜条件式','复制到第2门课程；再评估更多学校或硬件量产','主课程复购与服务成本稳定后另立预算','不预先承诺融资、论文录用或规模营收','项目负责人','T32','复用已验证的教材与数据流程']
]
rows(ws,28,roadmap)
ws['B28']='="M1｜"&TEXT($B$4,"m/d")&"—"&TEXT(EDATE($B$4,1)-1,"m/d")'
for r,label,m in [(29,'M2—M3',3),(30,'M4—M6',6),(31,'M7—M12',12),(32,'M13—M24',24)]:
    ws.cell(r,2,f'="{label}｜至"&TEXT(EDATE($B$4,{m})-1,"yyyy/m/d")')
    ws.cell(r,2).font=Font(name='微软雅黑',size=10,color='000000')
ws['B28'].font=Font(name='微软雅黑',size=10,color='000000')
band(ws,34,'D｜阅读依据：点击来源；具体指标的口径与来源另附于单元格批注')
refs = [
 ['用户确认','已完成基础','公司 / 4班约150人 / 微调 / 1000元','确认日：2026-09-17','本表7—13行','规划≠实绩','角色待绑定实名','项目组补齐原始凭证'],
 ['项目材料','课程与硬件','课程README、个性化实现、硬件STATUS','见单元格批注的本地路径','读取日：2026-09-17','以最新实绩更新','旧文档不覆盖新确认','8模块100知识点按仓库记载'],
 ['课程依据','Python官方教程','覆盖校对；本表不照搬教程顺序','https://docs.python.org/zh-cn/3/tutorial/','官方文档','用于教材设计','详见表02','不作为教学效果证明'],
 ['教学依据','Freeman et al. 2014','主动练习与即时反馈的设计参考','https://doi.org/10.1073/pnas.1319030111','PNAS研究','用于教学活动','详见表02','PLEX效果另行实测'],
 ['研究依据','WWC Handbook v5.0','基线、流失、班级聚类与研究解释','https://ies.ed.gov/ncee/wwc/handbooks','官方研究规范','用于研究方案','详见表03','阈值为本计划建议，非WWC认证'],
 ['模型依据','Model Cards','模型用途、评测条件与局限记录','https://research.google/pubs/model-cards-for-model-reporting/','研究原文','用于模型卡','详见表03','不等于模型性能已获认证']
]
rows(ws,35,refs)
for r in range(37,41):
    ws.cell(r,4).hyperlink = ws.cell(r,4).value
    ws.cell(r,4).font = Font(name='微软雅黑', size=9, color=TEAL, underline='single')
for r,k in [(37,'PY'),(38,'AL'),(39,'WWC'),(40,'MC')]: comment(ws.cell(r,2),k)
ws.row_breaks.append(Break(id=13))
ws.row_breaks.append(Break(id=24))

# 02: a teachable textbook, explicit personalization and real hardware use.
ws = sheet('02教材软硬件','PLEX｜教材、个性化与软硬件','课程范围：零基础大学生 Python 入门；建议8单元×4课时。课时与顺序待任课教师确认，现有100知识点先做映射与质量审核。',[11,23,37,40,37,35,32])
band(ws,4,'A｜8个教材单元：每单元都有可交付作品，而不是只有知识目录')
header(ws,5,['单元','教学主题','学生可完成的任务','数字教材交付','平台中的学习活动','硬件 / 数据连接','学习评价'])
units = [
 ['U1','环境、程序与调试','独立运行脚本，解释报错并修复缩进','环境指南、示范视频、调试工作纸','前测→短示例→代码修改→出口题','先用终端事件样例，不要求有设备','提交可运行脚本+修错说明'],
 ['U2','变量、类型与输入输出','完成类型转换、表达式与输入校验','概念卡、正反例、分层练习','按类型混淆错误推送补练','读取模拟按键数、倒计时数值','输入边界用例通过；解释转换原因'],
 ['U3','条件与布尔逻辑','编写理解度状态判断规则','流程图、分支追踪、错误诊断','预测输出→运行→解释差异','按“求助/完成/未响应”生成反馈','正确处理多条件、边界与异常值'],
 ['U4','循环与基础算法','汇总课堂事件，完成计数与搜索','循环可视化、追踪表、算法实验','补齐循环不变量直觉与终止条件','统计终端事件，不访问真实身份','测试零条、多条、边界数据'],
 ['U5','字符串与文本处理','清洗题目文本、解析简单事件行','切片案例、字符串实验、纠错卡','同知识点三档任务，逐步撤除提示','处理设备匿名事件文本','不丢有效字段；能说明清洗规则'],
 ['U6','列表、字典与数据组织','用列表/字典组织记录并去重汇总','数据结构图、练习集、项目模板','比较两种实现并解释选型','按匿名会话ID汇总求助事件','结果正确，能处理重复与缺失'],
 ['U7','函数、模块与测试','把重复逻辑封装为可复用函数','函数契约、模块示例、测试清单','先补全函数，再独立设计接口','Python脚本调用测试API或读CSV','函数职责清晰；正常/异常用例通过'],
 ['U8','文件、异常与综合实践','完成学习数据报告与课堂求助助手','2个项目说明、评分量规、反思模板','项目制：计划→编码→测试→互评','用真实脱敏事件与模拟事件分别验证','独立迁移任务；代码、报告和反思齐全']
]
rows(ws,6,units)
comment(ws['B6'],'PY','K','PLAN')
note(ws,14,'统一单元结构：学习目标→诊断→讲解与正反例→动手练习→形成性评价→错因反馈→迁移任务；评分建议：正确性50%、解释20%、调试20%、反思10%。')
band(ws,16,'B｜个性化生成依据：让每一次推荐都能回答“为什么给这个学生这份资源”')
header(ws,17,['链路','依据 / 输入','决策规则','产出与留痕','质量验收（建议）','教师 / 学生控制','任务关联'])
chain = [
 ['P1 诊断','前测、近期答题、用时、错因、学习目标','以掌握度和先修缺口为主；偏好仅调呈现','匿名ID、知识点、事件时间、画像版本','缺数据时标注低置信度并追加短测','学生可纠正画像；不固定贴“学习风格”标签','T02、T05'],
 ['P2 定位','先修图谱、薄弱点、复习到期与课时','先补前置，再定可完成的小目标','目标知识点、先修路径、难度与时长','抽样100条均可回溯到学习证据','教师可覆盖难度和任务目标','T05、T06'],
 ['P3 检索','审定教材、题库、代码例题与来源版本','限定课程范围，优先检索可引用证据','document_id、文档版本、片段与检索分数','独立检索集覆盖8单元；记录Recall@k','证据不足改用审定模板或转教师','T04—T06'],
 ['P4 生成','目标、画像快照、检索片段、约束','生成讲解、思维图、练习、拓展阅读或实验','资源ID、模型/提示词版本、参数、推荐理由','差异应体现在难度、支架和练习；不是只换称呼','保留标准教材入口；提供逐级提示','T05、T09'],
 ['P5 审核','生成资源、引用、答案与运行结果','检查事实、先修适配、答案与代码安全','自动检查、审核人、通过/退回与原因','抽样100条教师合格率≥90%；关键代码全部可运行','低置信度和高风险内容必须教师复核','T04、T06、T19'],
 ['P6 使用','已批准资源与当前教学节奏','分块呈现、先作答再反馈、错后提示','展示→接受→学习→作答的完整事件链','推荐接受率与学习完成率分别统计','学生可跳过；核心教材可低带宽访问','T05、T29'],
 ['P7 更新','退出题、迁移题、复习结果与教师反馈','更新掌握度、安排间隔复习，再生成下一步','同一trace_id关联画像—资源—评价—新画像','评价不重复原题；不将点击或灯色等同掌握','确认后的目标不被单次行为静默覆盖','T05、T11、T12']
]
rows(ws,18,chain)
for r in range(18,25): comment(ws.cell(r,3),'P','PLAN')
comment(ws['C18'],'PLAN','AL')
band(ws,26,'C｜软硬件与交付：共享课堂终端 + Python数据实践')
header(ws,27,['工作包','场景 / 范围','具体实现','学生与教师收益','验收 / 放行条件','成本与限制','任务关联'])
hardware = [
 ['H1','实物与必要性核查','盘点现有样机、固件和设备API；比较网页按钮','明确终端能解决的课堂问题','M1形成最新状态表、照片与差异清单','旧STATUS为设计版；若已有实物则直接补验收','T14'],
 ['H2','1套端云闭环','ESP32-S3+屏幕+4按键；匿名会话→事件API→看板','任务提醒、理解反馈、求助与教师响应','真实设备演示；事件不重计；断网可恢复','固件可为C++；Python课通过CSV/API学习，不强行烧录','T15'],
 ['H3','可靠性与小批量','72小时运行；断网、令牌失效、重复上报测试','减少课堂中断并保护学习记录','事件送达率≥99%；先10人易用性，再最多10套','234元/套仅旧BOM估算，另计外壳、维护、检测','T16—T17'],
 ['H4','教学增益验证','相同内容下比较软件入口与硬件入口','检验求助响应、注意干扰与参与成本','记录求助响应中位数、任务完成率与教师操作时长','硬件独立评价；不要把综合干预效果全归因硬件','T18'],
 ['D1','教师可独立上课','课前诊断→发布→课堂干预→课后复习→导出','把平台纳入一节完整课与一门完整课','1名未参与开发教师按手册独立完成1轮','教师端保留审核和覆盖权；培训用时计入成本','T29—T30'],
 ['D2','教材维护与可访问','作者、来源、许可、勘误、版本；字幕与文本替代','可持续更新，设备不足仍可完成学习','每单元有来源与变更记录；键盘可用、低带宽可读','AI生成内容经过审核再发布；禁用无授权素材','T03—T04、T29']
]
rows(ws,28,hardware)
comment(ws['C28'],'H')
comment(ws['F30'],'B')
ws.row_breaks.append(Break(id=14))
ws.row_breaks.append(Break(id=24))

# 03: existing work is preserved, conclusions remain tied to evidence.
ws = sheet('03科研模型','PLEX｜科研与模型证据闭环','已有基础：4班约150人测试、数据处理与模型微调已完成。下一步是复核、复现、解释和独立验证；不预设显著提升或论文录用。',[12,25,44,42,42,34,19])
band(ws,4,'A｜研究设计与数据：先分清“已有发现”与“下一轮要验证的假设”')
header(ws,5,['编号','要解决的问题','具体做法','口径 / 成功条件','必须保留的证据','结论边界','任务关联'])
research = [
 ['R01','已有150人的实绩复盘','还原4班测试时间、版本、干预方式和对照条件','约150人先去重；分别报告参与/有效配对/退出人数','匿名名单、原始导出、教师证明、测验版本','已完成研究可回顾登记；不得补写为事前预注册','T01、T07'],
 ['R02','明确教育研究问题','主问题：个性化资源是否改善Python迁移能力；次问题：是否节省教师时间','主指标只设1项；新一轮可设6周干预+2周延迟测','分析方案、课程目标、评分量规、时间表','时长为建议；先对齐学期，不承诺正向效果','T10'],
 ['R03','选择可信对照','同内容/时长/教师支持比较常规资源与个性化资源；条件允许再随机分配','记录分配层级、基线成绩与教师差异','分组日志、实际分配过程、执行偏差','只有4班时为探索性；不把150人当150个独立班级','T10—T11'],
 ['R04','决定样本与新增队列','用既有方差、班内相关和最小有意义差异做功效或精度分析','300人是运营累计目标；科研样本由设计决定','样本量脚本、假设、目标班数和流失预留','必要时增加班数；只增每班人数不解决聚类不足','T10、T28'],
 ['R05','保持评价独立','前/后/延迟测用等值题；迁移题不进训练或提示词','盲评双人抽评至少20%；报告评分一致性','题卷、量规、评分记录、评分分歧处理','训练内题成绩只能表示练习表现','T02、T10—T11'],
 ['R06','数据清洗与授权','建匿名ID/班级/版本/资源/结果字典；去重、缺失、异常处理可重跑','教学使用、研究与训练用途分别告知和记录授权依据','数据字典、清洗脚本、排除日志、用途与保留期记录','脱敏不等于匿名化；不上传可识别学生信息给外部模型','T07、T20'],
 ['R07','报告可复核统计','估计基线调整后的组间差异、效应量与95%区间；适配班级分配层级','报告主指标、退出/缺失和敏感性分析；副指标注明探索性','冻结数据哈希、分析代码、图表与完整报告','班级少时区间不稳，强调探索性，不强作因果结论','T12—T13'],
 ['R08','控制解释范围','分别检验学习结果、教师负担、用户体验和硬件价值','现有结论逐条关联对应表格或图；无效结果也保留','结论—证据映射、反例、版本适用范围','满意度/点击率不替代学习增益；整体效果不等于某Agent功劳','T12—T13、T18']
]
rows(ws,6,research)
for r in [8,9,10,12]: comment(ws.cell(r,3),'WWC','PLAN')
comment(ws['C11'],'LAW','PLAN')
band(ws,15,'B｜微调完善：用更干净的数据和同条件评测，证明哪些改动值得保留')
header(ws,16,['编号','工作环节','具体做法','口径 / 放行条件','必须保留的证据','失效时怎么处理','任务关联'])
model = [
 ['M01','已有微调可复现','登记基座与许可、微调方法、参数量、超参、随机种子和训练环境','用同一数据版本复跑；记录差异与算力费用','训练配置、日志、数据哈希、权重/适配器、代码提交','缺哪项补哪项，不把调用API表述为自研基座','T09'],
 ['M02','防止测试泄漏','核查现有150人数据是否参与训练或调参；按学生/时间/题目近重复隔离','训练/验证/测试名单冻结；用过的数据不再冒充独立测试','去重报告、分组清单、污染检查与新队列测试集','如现有测试已泄漏，重建新学期或新班级测试集','T08、T10'],
 ['M03','同条件消融比较','A基座；B基座+RAG；C微调+相同RAG；再按需比较画像/多Agent','固定题集、生成参数和预算；至少200条覆盖8单元的独立任务','模型输出、盲评量规、成本、耗时、失败样本','先确认RAG与简单流程价值，再判断微调/多Agent增量','T06、T09'],
 ['M04','评测教育质量与安全','评准确性、引用支持、难度适配、提示质量、代码通过率和泄漏风险','建议盲评合格率≥90%；重大事实/隐私错误0；报告样本分母','双人复核至少20%；保留分歧、置信区间与失败类别','未过门槛的版本不替换线上模型；严重错误先修复','T06、T09、T19'],
 ['M05','模型上线与回滚','模型卡、版本注册、少量灰度、成本告警、教师接管','预先冻结主质量指标的最小有意义改进；成本在预算内','版本对照、每任务费用、延迟分位数、回滚演练','未见可靠增益则保留基线；微调成果可作为负结果报告','T09、T19'],
 ['M06','数据飞轮可控更新','仅用授权且教师复核的困难样本追加训练；每轮保留旧版基准','训练集增长有质量记录；独立测试不反复调参消耗','样本来源、筛选规则、删除请求记录、版本变更','漂移或关键分组性能恶化时暂停更新','T08—T09、T20'],
 ['M07','科研与产品成果交付','1份研究报告/论文初稿；1张模型卡；1套可复现代码与脱敏样例','指导教师复核；结果、局限、数据可用性声明齐全','报告、模型卡、统计脚本、审阅记录','课程结题不以论文录用、专利授权或正结果为前提','T13、T30']
]
rows(ws,17,model)
for r in range(17,24): comment(ws.cell(r,3),'PLAN')
comment(ws['C21'],'MC','PLAN')
note(ws,25,'核心口径：150名学生 ≠ 150条训练样本 ≠ 150次使用；课程效果 ≠ 模型离线得分；微调完成 ≠ 微调有效。每项结论同时注明对象、版本、时间、分母和证据。')
ws.row_breaks.append(Break(id=13))

# 04: transparent and editable scenario model, no fabricated actual revenue.
ws = sheet('04商业公益','PLEX｜商业模式与公益测算','已有营收1000元为用户确认；下方未来12个月全部为情景假设，非实际价格或订单。蓝字黄底可改；公式自动联动。金额均为人民币。',[32,19,19,19,17,22,23,26])
band(ws,4,'A｜商业闭环：已有公司做规范化经营，先服务好一门课程')
for a,b,t in [(1,2,'产品 / 付费方'),(3,4,'交付与收费'),(5,6,'获客与复购'),(7,8,'履约与边界')]:
    ws.merge_cells(start_row=5,start_column=a,end_row=5,end_column=b)
    ws.cell(5,a,t).font = Font(name='微软雅黑',size=10,bold=True,color='FFFFFF')
    ws.cell(5,a).fill = PatternFill('solid',fgColor=NAVY)
    ws.cell(5,a).alignment = Alignment(vertical='center')
offers = [
 ['学生课程包','一学期教材+分层练习+限额辅导；基础教材有免费入口','教师转介、公开体验课→激活→付费→完成→续购','明确期限、额度与退款规则；不付费不影响实验参与及成绩'],
 ['机构班级包','学校/培训方支付：班级空间、教师看板、资源与服务','现有4班案例→教师演示→机构试用→合同→复购','按机构包核算席位；同一服务不再向覆盖学生重复收费'],
 ['硬件服务','先借用/租用共享终端；采购另立报价和交付条件','以减少求助等待或课堂操作负担证明价值','硬件销售收入不纳入本模型；未验收不计付费产品'],
 ['学生半价 / 助学','核验学生资格后课程包5折；另设免费助学名额','公开规则、名额、期限和受益人数；教师协助触达','与研究授权分离；价格让利≠现金捐赠；免费名额不计营收']
]
for r, data in enumerate(offers,6):
    for (a,b), txt in zip([(1,2),(3,4),(5,6),(7,8)],data):
        ws.merge_cells(start_row=r,start_column=a,end_row=r,end_column=b)
        c=ws.cell(r,a,txt); c.font=Font(name='微软雅黑',size=10,color=INK)
        c.alignment=Alignment(wrap_text=True,vertical='center')
        c.fill=PatternFill('solid',fgColor='F2F6FA' if r%2==0 else 'FFFFFF')
    ws.row_dimensions[r].height=56
comment(ws['G6'],'PAY','PLAN')
band(ws,11,'B｜未来12个月投入与销量假设：三种情景均可编辑，不混入历史收入')
header(ws,12,['参数','保守','基准','进取','单位','口径 / 来源','',''])
ws.merge_cells('F12:H12')
assumptions = [
 (13,'测算周期',12,12,12,'月','未来12个月；不把全年销量再乘12'),
 (14,'学生课程包原价',99,99,99,'元/人期','规划试价；先做访谈与小额真实付费验证'),
 (15,'学生半价折扣',.5,.5,.5,'比例','按同一服务公示原价5折；不虚构划线价'),
 (16,'全价课程包销量',30,60,100,'人期','与半价销量互斥；销量不等于独立付费人数'),
 (17,'半价课程包销量',20,40,80,'人期','真实付费人期；期内按单去重'),
 (18,'机构包销量',1,2,3,'单','均为未来新增已履约场景假设'),
 (19,'机构包单价',1500,2000,2500,'元/单','独立合同定价假设；尚非签约金额'),
 (20,'每机构包服务席位',40,40,40,'人期/单','机构覆盖者不重复计算为个人付费销量'),
 (21,'预计退款比例',.05,.05,.05,'比例','基于收费总额估计；实际改用订单退款记录'),
 (22,'支付手续费率',.006,.006,.006,'比例','示例费率，待合同替换；按退款前收款计费'),
 (23,'直接服务变动成本',12,12,15,'元/人期','含AI调用/可变运维与履约；所有席位和免费名额均计'),
 (24,'月固定现金成本',500,500,800,'元/月','云资源、基础工具与固定服务费用；团队薪酬见等价工时'),
 (25,'新增硬件数量',1,10,10,'套','先验收1套，再扩至最多10套；存量不重复采购'),
 (26,'单套BOM估算',234,234,234,'元/套','仓库旧BOM估算；非采购报价，不含人工/开模/检测/税费'),
 (27,'硬件追加费用',766,1660,2660,'元','外壳、维护、检测等规划预留；需采购前更新'),
 (28,'新增训练与评测预算',1000,2000,4000,'元','未来新增支出；已完成微调历史成本另行归档'),
 (29,'活动与获客现金预算',500,1000,2000,'元','宣传、活动、差旅等；不重复计价格让利'),
 (30,'免费助学名额',10,20,30,'人期','与全价/半价/机构服务不重叠；仍发生服务成本'),
 (31,'未付费团队与志愿工时',240,480,720,'小时','含开发、教研、运营与活动；如有工资应调整避免重复'),
 (32,'工时机会成本单价',40,40,40,'元/小时','仅作完整成本估算，非实际工资或现金支付'),
 (33,'历史已实现营收',"='01目标路线'!B9","='01目标路线'!B9","='01目标路线'!B9",'元','用户确认1000元；与净收款合计前须统一退款口径'),
 (34,'累计净收款规划目标',10000,10000,10000,'元','本计划建议；用于看差额，不是收入预测保证')
]
for r,name,a,b,c,unit,desc in assumptions:
    rows(ws,r,[[name,a,b,c,unit,desc,None,None]],32)
    ws.merge_cells(start_row=r,start_column=6,end_row=r,end_column=8)
    ws.row_dimensions[r].height=34
    for col in range(2,5):
        cell=ws.cell(r,col)
        cell.number_format=PCT if r in [15,21,22] else (MONEY if r in [14,19,23,24,26,27,28,29,32,33,34] else NUM)
        if r!=33:
            cell.fill=PatternFill('solid',fgColor=YELLOW); cell.font=Font(name='微软雅黑',size=11,color=BLUE)
            comment(cell,'B' if r==26 else 'PLAN')
        else: comment(cell,'U')
    for col in range(1,6): ws.cell(r,col).alignment=Alignment(wrap_text=True,vertical='center')
dv = DataValidation(type='decimal',operator='between',formula1=0,formula2=1)
dv.errorTitle='比例输入错误'; dv.error='请输入0到1之间的比例。'; dv.showErrorMessage=True
ws.add_data_validation(dv)
for r in [15,21,22]: dv.add(f'B{r}:D{r}')
nonneg = DataValidation(type='decimal',operator='greaterThanOrEqual',formula1=0)
nonneg.showErrorMessage=True; nonneg.error='请输入大于或等于0的数。'; ws.add_data_validation(nonneg)
for r in [13,14,16,17,18,19,20,23,24,25,26,27,28,29,30,31,32,34]: nonneg.add(f'B{r}:D{r}')
band(ws,36,'C｜自动测算：基准情景若亏损，先优化成本和客单价，不用规模掩盖问题')
header(ws,37,['指标','保守','基准','进取','单位','计算口径','',''])
ws.merge_cells('F37:H37')
metrics = [
 (38,'个人付费销量','SUM({c}16:{c}17)','人期','全价+半价；独立付费人数须另按用户去重'),
 (39,'全部服务量','{c}38+{c}18*{c}20+{c}30','人期','个人+机构席位+免费；三类互斥假设'),
 (40,'个人课程包收款','{c}14*{c}16+{c}14*{c}15*{c}17','元','优惠按实际成交价计入'),
 (41,'机构包收款','{c}18*{c}19','元','机构包销量×单价'),
 (42,'新增收费总额','SUM({c}40:{c}41)','元','仅未来新增，不含历史1000元'),
 (43,'预计退款','{c}42*{c}21','元','交易总额×退款率'),
 (44,'新增净收款','{c}42-{c}43','元','收款减退款；非会计确认营收'),
 (45,'支付手续费','{c}42*{c}22','元','按总收款计费的保守假设'),
 (46,'变动履约成本','{c}39*{c}23','元','包含免费名额成本；退款后服务成本保守保留'),
 (47,'经营贡献','{c}44-{c}45-{c}46','元','净收款减支付费和直接履约成本'),
 (48,'经营贡献率','IFERROR({c}47/{c}44,0)','比例','净收款为0时显示—'),
 (49,'固定现金支出','{c}13*{c}24','元','12个月固定成本'),
 (50,'硬件现金投入','{c}25*{c}26+{c}27','元','新增数量×BOM+追加费用；现金口径一次列入'),
 (51,'训练及获客支出','{c}28+{c}29','元','只计未来新增预算'),
 (52,'预计现金结余','{c}47-SUM({c}49:{c}51)','元','不含历史现金；未计税费，非净利润'),
 (53,'团队工时等价成本','{c}31*{c}32','元','无薪工作也有成本；已付薪时调整避免重复'),
 (54,'含工时成本结余','{c}52-{c}53','元','体现学生团队投入；不是会计利润'),
 (55,'半价让利金额','{c}14*(1-{c}15)*{c}17','元','相对公示原价的优惠；不是额外现金支出'),
 (56,'免费名额标价价值','{c}14*{c}30','元','对外需注明标价口径；不是现金捐赠'),
 (57,'助学直接履约成本','({c}17+{c}30)*{c}23','元','半价+免费人期的服务成本；已含在变动成本中'),
 (58,'累计口径对齐后金额','{c}33+{c}44','元','仅在历史1000元也为净收款时可比较'),
 (59,'距累计目标差额','MAX(0,{c}34-{c}58)','元','不足部分；达到目标显示—'),
 (60,'个人单人期边际贡献','IFERROR(({c}40*(1-{c}21-{c}22))/{c}38-{c}23,0)','元/人期','保持原全价/半价结构；扣退款、手续费和变动成本'),
 (61,'现金保本还需个人销量','IF({c}60>0,ROUNDUP(MAX(0,-{c}52)/{c}60,0),"无法估算")','人期','固定成本不增、原折扣结构不变的近似增量；非销售承诺')
]
for r,name,expr,unit,desc in metrics:
    rows(ws,r,[[name]+['='+expr.format(c=c) for c in 'BCD']+[unit,desc,None,None]],32)
    ws.merge_cells(start_row=r,start_column=6,end_row=r,end_column=8)
    ws.row_dimensions[r].height=34
    for col in range(2,5):
        cell=ws.cell(r,col); cell.number_format=PCT if r==48 else (NUM if r in [38,39,61] else MONEY)
        if r in [44,47,52,54,58,59,61]: cell.fill=PatternFill('solid',fgColor=LIGHT); cell.font=Font(name='微软雅黑',size=11,bold=True,color='000000')
    comment(ws.cell(r,1),'PLAN')
ws.conditional_formatting.add('B52:D54',FormulaRule(formula=['B52<0'],fill=PatternFill('solid',fgColor='FCE4D6'),font=Font(color='A33224')))
note(ws,63,'执行口径：逐笔对齐订单—支付—退款—权益—履约—票据；月度记录转化、续购、投诉、单人成本。研究参与、评分和必要教学资源不与购买或训练授权绑定。')
note(ws,64,'助学口径：分别公开半价人数、免费人数、实际付费、价格让利与直接服务成本；不把优惠包装为现金捐赠。公司、合同、知识产权、税务和支付规则由负责人按实际业务复核。',fill=LIGHT)
comment(ws['A63'],'PAY','LAW','PLAN')
comment(ws['A64'],'PLAN')
ws.freeze_panes='B4'

# 05: a bounded and assignable execution register.
ws = sheet('05执行证据','PLEX｜执行与证据清单','32项增量任务；M1从表01启动日算起，日期自动联动。负责人待替换实名；依赖项须在交付验收前完成，可并行准备。状态和证据入口可编辑。',[10,10,36,19,16,16,24,48,15,30])
note(ws,4,'首周优先：T01固化已有成果、T07整理原始数据、T14核实硬件实况、T20确认数据用途、T22核对公司与账务。以下“待开始”仅指完善任务，不否定已完成成果。',fill=LIGHT)
header(ws,5,['ID','优先级','下一步动作','主责角色','开始日期','截止日期','交付依赖','交付物 / 验收标准','任务状态','证据入口（回填）'])
tasks = [
 ('T01','P0','固化已有成果与最新项目说明','项目负责人',0,1,'—','公司、4班约150人、微调、1000元分别有证据索引；更新旧文档'),
 ('T02','P0','冻结课程目标、8单元矩阵与评分量规','教研负责人',0,1,'T01','100知识点映射到8单元；前后测及2项目量规经教师确认'),
 ('T03','P0','把课程库编成可授课教材V1','教研 / 内容',1,2,'T02','8单元教材、例题、分层练习、教师教案与勘误入口'),
 ('T04','P0','审查教材、答案与版权来源','教研负责人',2,3,'T03','100知识点双人审稿；关键示例可运行；素材来源可查'),
 ('T05','P0','补全个性化依据、追踪与反馈字段','产品 / 后端',1,2,'T02、T20','画像版本—检索引用—生成版本—审核—学习结果能一链追溯'),
 ('T06','P0','建立真实输出与检索质量评测','算法 / 教研',2,3,'T04、T05、T08','100条资源人工审查；200条独立模型任务；保留失败样例'),
 ('T07','P0','复盘已完成4班测试并清洗数据','科研负责人',0,1,'—','参与/去重/配对/退出人数、字典、清洗脚本与排除日志'),
 ('T08','P0','核查训练授权和测试泄漏','算法 / 科研',1,2,'T07、T20','学生/时间/近重复题目隔离；训练、验证、测试集版本冻结'),
 ('T09','P0','复现现有微调并产出模型卡','算法负责人',2,3,'T08','基座/RAG/微调同条件比较；质量、成本、局限和回滚版本齐全'),
 ('T10','P1','锁定新一轮实验与教师合作','科研 / 教研',3,4,'T02、T07、T08、T09','预先登记主指标、分组、样本量依据、知情流程与时间窗'),
 ('T11','P1','执行独立新队列教学验证','教研 / 科研',4,6,'T10、T19','按协议完成前测、干预、后测与延迟测；原始数据冻结'),
 ('T12','P1','形成统计、分组与缺失分析','科研负责人',6,7,'T11','报告基线调整效应、95%区间、退出和敏感性分析'),
 ('T13','P1','交付科研报告与论文初稿','科研负责人',7,8,'T12、T18','1份可复核报告/初稿；结论对应证据；无效与局限如实呈现'),
 ('T14','P0','核实硬件进展及课堂必要性','硬件 / 教研',0,1,'—','最新实物/固件/API状态表；说明相对网页的预期增益'),
 ('T15','P1','完成1套真实硬件端云联通','硬件 / 后端',1,2,'T14、T20','按键→事件→教师看板→响应；附设备、固件与演示记录'),
 ('T16','P1','验收可靠性与小样本易用性','测试 / 硬件',2,3,'T15','72小时、断网/重连/重复上报；10人易用性；失败列表关闭'),
 ('T17','P1','条件通过后制作最多10套试点终端','硬件 / 运营',3,5,'T16、T22','采购与资产台账、维护手册；真实照片与成本凭证'),
 ('T18','P1','单独验证硬件的课堂价值','教研 / 科研',5,6,'T17','比较求助响应、课堂操作与干扰；有增益才继续投入'),
 ('T19','P0','完成正式交付与恢复演练','工程负责人',2,3,'T05、T09、T20','权限、代码执行隔离、备份恢复、失败降级与回滚；4周可用率≥99%'),
 ('T20','P0','确认数据用途、权限与删除流程','数据负责人',0,1,'—','用途、授权依据、存储位置、访问人、保留期、退出/删除流程可查'),
 ('T21','P0','审计现有收费全链路并补齐缺口','后端 / 财务',1,2,'T01、T22','下单、回调验签/幂等、退款、权益回收、对账与票据真实走通'),
 ('T22','P0','整理公司、合同与财务基础','财务 / 项目',0,1,'—','主体与账户、历史1000元对账、教材/代码/数据权属、服务合同归档'),
 ('T23','P1','校准产品包、定价与服务成本','运营 / 财务',1,2,'T01、T22','访谈至少2名机构决策者和10名学生；更新表04价格/成本假设'),
 ('T24','P1','跑通稳定付费转化与履约','运营负责人',3,6,'T04、T19、T21、T23','阶段目标：新增1机构单+50个人付费人期；记录渠道、退款与履约'),
 ('T25','P1','规范半价及免费助学规则','运营 / 教研',1,2,'T22、T23','资格、名额、期限、公示价格、预算及申诉入口；与实验参与分离'),
 ('T26','P1','形成半年度助学与运营台账','运营 / 财务',5,6,'T24、T25','受益人数去重；列实际收款、让利、免费名额、服务成本和反馈'),
 ('T27','P1','复制到第二个付费机构','运营 / 教研',8,9,'T13、T24、T29','第二个真实付费机构订单及教师独立使用记录；收款与履约匹配'),
 ('T28','P2','达成可控扩展并发布教材V2','项目 / 教研',9,12,'T13、T27','累计≥8班≥300去重学生；累计口径对齐后≥1万元；基准助学40半价+20免费'),
 ('T29','P1','形成教师可独立使用的交付包','工程 / 教研',3,5,'T04、T19','新环境部署、教师手册、培训与支持；非开发教师独立完成1轮'),
 ('T30','P0','提交课程项目完整成果包','项目负责人',2,3,'T04、T06、T09、T16、T19、T21','教材样章+完整课例+运行演示+已有实证+模型卡+商业/公益台账'),
 ('T31','P2','年度经营与全成本复盘','财务 / 项目',11,12,'T26、T27、T28','月度收入/退款/成本可对账；连续3月经营贡献为正；如实列完整成本'),
 ('T32','P2','决定第2课程与下一年度投入','项目负责人',11,12,'T28、T31','仅在复购、教学质量与成本通过评审后立项；形成下一年预算')
]
for r,t in enumerate(tasks,6):
    ident,pri,action,owner,start,end,dep,accept=t
    rows(ws,r,[[ident,pri,action,owner,f"=EDATE('01目标路线'!$B$4,{start})",f"=EDATE('01目标路线'!$B$4,{end})-1",dep,accept,'待开始','待归档']],62)
    ws.cell(r,5).number_format=DATE; ws.cell(r,6).number_format=DATE
    for c in [4,9,10]:
        ws.cell(r,c).fill=PatternFill('solid',fgColor=YELLOW)
        ws.cell(r,c).font=Font(name='微软雅黑',size=10,color=BLUE)
    comment(ws.cell(r,8),'PLAN')
    if pri=='P0':
        ws.cell(r,2).font=Font(name='微软雅黑',size=11,bold=True,color='A34A20')
        ws.cell(r,2).fill=PatternFill('solid',fgColor='FCE4D6')
status = DataValidation(type='list',formula1='"待开始,进行中,待验收,已完成,阻塞"')
status.showErrorMessage=True; status.error='请选择列表中的任务状态。'; ws.add_data_validation(status); status.add('I6:I37')
for state,color in [('已完成','DFF0E5'),('进行中','E0F0FC'),('待验收','FFF2CC'),('阻塞','FCE4D6')]:
    ws.conditional_formatting.add('I6:I37',FormulaRule(formula=[f'I6="{state}"'],fill=PatternFill('solid',fgColor=color)))
table=Table(displayName='PlexExecution',ref='A5:J37')
table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showFirstColumn=False,showLastColumn=False,showRowStripes=False,showColumnStripes=False)
ws.add_table(table)
ws.freeze_panes='D6'
ws.print_title_rows='1:5'

for ws in wb:
    ws.sheet_properties.outlinePr.summaryRight=False
    ws.print_area=f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'
    ws.sheet_view.tabSelected=False
wb.active=0
wb['01目标路线'].sheet_view.tabSelected=True
wb.calculation.fullCalcOnLoad=True
wb.calculation.forceFullCalc=True
wb.properties.title='PLEX 长期完善与落地计划'
wb.properties.subject='数字教材开发课程：教学、科研、技术、商业与公益'
wb.properties.creator='PLEX 项目规划'
wb.properties.description='基于用户2026-09-17确认与指定本地项目文档；五张工作表，规划假设与既有成果分开记录。'
wb.save(FILE)
print(FILE)
print({'sheets':wb.sheetnames,'tasks':len(tasks),'formulas':sum(c.data_type=='f' for s in wb for row in s for c in row)})
