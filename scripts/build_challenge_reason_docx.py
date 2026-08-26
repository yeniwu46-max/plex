from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "submission" / "artifacts" / "挑战赛理由及作品亮点.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(24, 36, 52)
MUTED = RGBColor(92, 104, 119)
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
BORDER = "D9E2EF"


def set_run_font(run, name="Microsoft YaHei", size=None, color=None, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def style_paragraph(paragraph, before=0, after=6, line=1.10, align=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), size)
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)


def set_table_width(table, widths):
    table.autofit = False
    table.allow_autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_para(doc, text="", bold_prefix=None):
    p = doc.add_paragraph()
    style_paragraph(p, after=6, line=1.10)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True, color=INK)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, color=INK)
    else:
        r = p.add_run(text)
        set_run_font(r, color=INK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading("", level=level)
    if level == 1:
        style_paragraph(p, before=16, after=8)
        size, color = 16, BLUE
    elif level == 2:
        style_paragraph(p, before=12, after=6)
        size, color = 13, BLUE
    else:
        style_paragraph(p, before=8, after=4)
        size, color = 12, DARK_BLUE
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=True)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    style_paragraph(p, after=4, line=1.167)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    r = p.add_run(text)
    set_run_font(r, color=INK)
    return p


def add_metadata_rows(doc):
    rows = [
        ("作品名称", "PLEX Universe / 基于大模型的个性化资源生成与学习多智能体系统"),
        ("赛题方向", "A3 高等教育个性化学习资源多智能体系统"),
        ("提交类型", "第十五届“中国软件杯”大学生软件设计大赛挑战赛优化作品"),
        ("版本日期", "2026 年 7 月 31 日"),
    ]
    for label, value in rows:
        p = doc.add_paragraph()
        style_paragraph(p, after=2, line=1.10)
        r1 = p.add_run(f"{label}：")
        set_run_font(r1, bold=True, color=INK)
        r2 = p.add_run(value)
        set_run_font(r2, color=INK)


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    set_table_borders(table, color="C8D4E3", size="8")
    set_table_width(table, (6.5,))
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    tr = p.add_run(title)
    set_run_font(tr, bold=True, color=DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    br = p2.add_run(body)
    set_run_font(br, color=INK)
    doc.add_paragraph()


def add_score_table(doc):
    headers = ["评分项", "作品对应亮点"]
    rows = [
        (
            "创新价值与实用性",
            "将对话画像、多智能体资源生成、教师审核、学习路径和效果评估整合成高校编程课程的完整学习闭环，解决资源匹配和个性化指导不足的问题。",
        ),
        (
            "功能实现及技术要求",
            "已具备学生、教师、管理三端；支持画像、资源任务、试炼、星轨路径、班级观察、权限控制、统一 API、JWT 认证和可切换数据库基础。",
        ),
        (
            "配套文档丰富度",
            "已整理需求说明、系统设计、多智能体设计、知识库与评测、测试说明、部署说明、用户手册、开源与 AI 工具说明等配套材料。",
        ),
        (
            "演示视频与 PPT 效果",
            "挑战赛版本重点补齐演示材料，建议以“画像构建 -> 资源生成 -> 教师审核 -> 路径推送 -> 试炼反馈 -> 效果评估”的故事线呈现。",
        ),
    ]
    table = doc.add_table(rows=1, cols=2)
    set_table_borders(table)
    set_table_width(table, (1.75, 4.75))
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_FILL)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, bold=True, color=DARK_BLUE, size=10.5)
    for label, detail in rows:
        cells = table.add_row().cells
        for idx, value in enumerate((label, detail)):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_run_font(r, color=INK, size=10.2)
    doc.add_paragraph()


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for side in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, side, Inches(1.0))
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)

    for name in ("Heading 1", "Heading 2", "Heading 3"):
        styles[name].font.name = "Microsoft YaHei"
        styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("PLEX Universe 挑战赛提交材料")
    set_run_font(run, size=9, color=MUTED)
    return doc


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = setup_document()

    p = doc.add_paragraph()
    style_paragraph(p, before=16, after=4, line=1.10)
    r = p.add_run("参加挑战赛理由及作品亮点")
    set_run_font(r, size=22, color=RGBColor(0, 0, 0), bold=True)

    p = doc.add_paragraph()
    style_paragraph(p, after=14)
    r = p.add_run("面向 A3 高等教育个性化学习资源多智能体系统赛题的挑战赛优化说明")
    set_run_font(r, size=12.5, color=MUTED)

    add_metadata_rows(doc)
    add_callout(
        doc,
        "核心判断",
        "PLEX 的价值不在于单次生成学习资料，而在于把学生画像、多智能体资源生产、教师审核、路径推送、试炼反馈和效果评估串成可运行的教育闭环。",
    )

    add_heading(doc, "一、参加挑战赛理由")
    add_heading(doc, "1. 赛题方向与作品定位高度一致", level=2)
    add_para(
        doc,
        "A3 赛题要求参赛团队面向高等教育场景，利用大模型和多智能体技术解决“学习资源多而杂、学生差异大、个性化指导不足、学习效果难量化”的问题。PLEX 的产品目标正是把传统教学平台从“资源陈列与任务发布”升级为“画像驱动、资源自动生成、路径动态调整、教师可审核”的智能学习系统。",
    )
    add_para(
        doc,
        "作品已围绕学生端、教师端和管理端形成三类工作台：学生完成每日委托、试炼与个性化资源学习；教师查看班级热力图、学生档案、试炼表现与资源审核；管理员维护平台策略、AI 策略、通知开关和安全配置。这一结构与赛题中“学生个体学习需求 + 高校课程资源 + 多智能体协作”的业务场景具有直接对应关系。",
    )

    add_heading(doc, "2. 挑战赛是补齐初赛短板、提升完整度的关键机会", level=2)
    add_para(
        doc,
        "初赛版本已经具备用户认证、三端路由、学习任务、试炼、成就激励、画像接口、资源生成任务和教师审核等工程基础，但对照挑战赛提交要求与评分项，仍需要把若干能力从“已有后端或文档说明”推进到“评委可以直接看到、直接操作、直接理解”的状态。",
    )
    add_para(
        doc,
        "本次参加挑战赛的核心目的不是推倒重来，而是在原有系统基础上集中优化四类内容：补齐 PPT、演示视频、报名表和说明文档等提交材料；强化流式输出、Markdown 渲染、多模态卡片等现代 AI 产品体验；完善画像对话、资源生成、路径推送、试炼反馈之间的前端走查链路；将本地确定性验收、外部模型接入边界和 AI Coding 工具使用情况写清楚。",
    )

    add_heading(doc, "3. 作品具备继续竞争的工程基础", level=2)
    add_para(
        doc,
        "PLEX 已经完成 Vue 3 + Vite + Naive UI + Pinia + Axios 前端工程，Flask + SQLAlchemy + JWT 后端工程，以及本地 SQLite / 生产 MySQL 的可切换数据基础。API 统一采用 /api/v1 前缀和 REST JSON 响应格式，核心页面围绕学生、教师、管理员三端隔离建设，能够支撑挑战赛要求的可运行作品提交。",
    )
    add_para(
        doc,
        "从功能侧看，系统已经具备学生首页、探索舱、今日委托、试炼参与、探索档案、星轨路径、教师领航总览、星域观测、Explorer 档案、试炼中枢和控制中枢等模块。挑战赛阶段可以在此基础上重点打磨个性化学习链路和演示叙事，而不需要从零搭建基础管理系统。",
    )

    add_heading(doc, "4. 参赛有助于展示“因材施教”的数字化落地价值", level=2)
    add_para(
        doc,
        "高等教育中的编程学习差异明显：学生的知识基础、学习目标、认知方式、易错点、实践经验和学习节奏并不相同。传统平台通常只能提供统一课程资源，难以及时回答“这个学生现在该学什么、为什么这样推荐、学完后是否真的提升”。",
    )
    add_para(
        doc,
        "PLEX 希望通过可解释画像、多智能体资源生产、教师审核和学习效果评估，把“因材施教”落到可操作的软件流程中。挑战赛提供了一个把教育业务价值、AI 技术融合和工程实现质量集中呈现的机会。",
    )

    add_heading(doc, "二、作品亮点")
    highlights = [
        (
            "对话式学生画像，支撑真正的个性化入口",
            "系统围绕学生画像构建个性化学习入口，画像维度覆盖知识基础、学习目标、认知风格、学习偏好、学历或专业背景、兴趣方向、易错点与学习进度等信息，满足赛题“不少于 6 个维度”的要求。画像既可以来自学生自然语言对话，也可以结合错题、试炼、每日委托和学习报告等行为证据持续更新。",
        ),
        (
            "多智能体协作生产资源，过程可追踪、结果可审核",
            "作品将资源生成拆解为画像解释、知识检索、教学设计、资源生成、质量审核和路径规划等角色。每个智能体消费上一步结构化结果，并记录任务状态、实际后端、耗时、失败原因和审核摘要，避免把复杂教育资源生成包装成不可解释的黑盒。",
        ),
        (
            "至少五类资源与多模态呈现，覆盖不同学习方式",
            "系统围绕 Python 程序设计基础知识库和学生画像生成差异化资源，重点覆盖课程讲解文档、知识点图谱或思维导图、分层练习题、拓展阅读材料、代码实践案例，并在挑战赛版本强化语音讲解、短视频脚本或视频卡片、答疑图解等多模态呈现方式。",
        ),
        (
            "学习路径、试炼与评价闭环完整",
            "PLEX 的学生端以探索舱、今日委托、星轨路径、试炼、探索档案组织学习过程。教师发布试炼后，学生可以在探索舱和今日委托中收到题目；补给站紧急任务会根据薄弱知识点生成练习；星轨路径覆盖语言基础、算法基础、动态规划、计算几何、图论、数据结构等学域，并支持知识点级导航。",
        ),
        (
            "三端协同，适合真实高校课程使用",
            "学生端关注学习任务、个性化资源、路径推进和激励反馈；教师端关注班级热力图、学生档案、风险学生、试炼发布和资源审核；管理端关注平台规则、AI 策略、通知、安全和数据配置。三端通过统一 API 和权限控制协作，既能服务学生个体学习，也能服务教师班级管理。",
        ),
        (
            "工程实现规范，提交材料边界清楚",
            "作品采用前后端分离架构，前端为 Vue 3 + Vite + Naive UI，后端为 Flask + SQLAlchemy，认证采用 JWT，API 返回统一 JSON 结构。本地开发可使用 SQLite，生产环境建议 MySQL。系统提供测试账号、种子数据脚本、启动脚本和接口文档，便于评委复现核心流程。",
        ),
    ]
    for title, body in highlights:
        add_bullet(doc, f"{title}：{body}")

    add_heading(doc, "三、与评分项的对应关系")
    add_score_table(doc)

    add_heading(doc, "四、建议答辩表达")
    add_para(
        doc,
        "本作品的核心不是简单调用大模型生成一段学习资料，而是把大模型能力放进真实教学流程：先理解学生，再由多智能体协作生成资源，再经过教师审核和路径规划推送给学生，最后用试炼、错题和学习效果数据反向更新画像。挑战赛版本将重点补齐初赛中演示材料、流式交互、多模态呈现和前端走查链路的不足，使作品的技术价值、教育价值和工程完整度能够被评委直观看到。",
    )

    doc.core_properties.title = "参加挑战赛理由及作品亮点"
    doc.core_properties.subject = "PLEX Universe A3 挑战赛提交材料"
    doc.core_properties.author = "PLEX 项目组"
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(path)
