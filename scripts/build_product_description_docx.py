from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("outputs") / "网站产品说明文档.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(24, 36, 52)
MUTED = RGBColor(92, 104, 119)
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
BORDER = "D9E2EF"
WHITE = "FFFFFF"


def set_run_font(run, name="Microsoft YaHei", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), size)
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)


def set_table_width(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.allow_autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def style_paragraph(paragraph, before=0, after=6, line=1.10, align=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def add_para(doc, text="", style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    style_paragraph(p)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True, color=INK)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, color=INK)
    else:
        r = p.add_run(text)
        set_run_font(r, color=INK)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    style_paragraph(p, after=4, line=1.167)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    r = p.add_run(text)
    set_run_font(r, color=INK)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    style_paragraph(p, after=4, line=1.167)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    r = p.add_run(text)
    set_run_font(r, color=INK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading("", level=level)
    if level == 1:
        style_paragraph(p, before=16, after=8)
        size, color = 16, BLUE
    elif level == 2:
        style_paragraph(p, before=12, after=6)
        size, color = 13, BLUE
    else:
        style_paragraph(p, before=8, after=4)
        size, color = 12, DARK_BLUE
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=True)
    return p


def add_key_value_table(doc, rows, widths=(1.55, 4.95)):
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_borders(table)
    set_table_width(table, widths)
    for i, (label, value) in enumerate(rows):
        label_cell, value_cell = table.rows[i].cells
        set_cell_shading(label_cell, LIGHT_FILL)
        for cell in (label_cell, value_cell):
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        lr = label_cell.paragraphs[0].add_run(label)
        set_run_font(lr, bold=True, color=DARK_BLUE)
        vr = value_cell.paragraphs[0].add_run(value)
        set_run_font(vr, color=INK)
    doc.add_paragraph()
    return table


def add_matrix_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_borders(table)
    set_table_width(table, widths)
    hdr = table.rows[0].cells
    for idx, text in enumerate(headers):
        set_cell_shading(hdr[idx], LIGHT_FILL)
        p = hdr[idx].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, bold=True, color=DARK_BLUE, size=10.5)
    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            set_run_font(r, color=INK, size=10.2)
    doc.add_paragraph()
    return table


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    set_table_borders(table, color="C8D4E3", size="8")
    set_table_width(table, (6.5,))
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    tr = p.add_run(title)
    set_run_font(tr, bold=True, color=DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    br = p2.add_run(body)
    set_run_font(br, color=INK)
    doc.add_paragraph()


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    header = section.header.paragraphs[0]
    style_paragraph(header, after=0)
    hr = header.add_run("PLEX Universe / A3 个性化学习系统 | 网站产品说明文档")
    set_run_font(hr, size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    style_paragraph(footer, after=0)
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = footer.add_run("第 ")
    set_run_font(fr, size=9, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    footer.runs[-1]._r.append(fld_begin)
    footer.runs[-1]._r.append(instr)
    footer.runs[-1]._r.append(fld_end)
    rr = footer.add_run(" 页")
    set_run_font(rr, size=9, color=MUTED)
    return doc


def build():
    doc = setup_document()

    # Cover / masthead
    for _ in range(2):
        doc.add_paragraph()
    title = doc.add_paragraph()
    style_paragraph(title, before=0, after=4)
    r = title.add_run("网站产品说明文档")
    set_run_font(r, size=25, color=RGBColor(0, 0, 0), bold=True)
    subtitle = doc.add_paragraph()
    style_paragraph(subtitle, after=16)
    sr = subtitle.add_run("PLEX Universe / A3 个性化学习系统")
    set_run_font(sr, size=14, color=MUTED)
    add_key_value_table(
        doc,
        [
            ("文档定位", "第一部分：网站产品说明，用于课程/项目报告中的产品介绍与系统说明。"),
            ("产品类型", "面向编程学习场景的个性化学习平台，覆盖学生端、教师端、管理端。"),
            ("技术架构", "前后端分离：Vue 3 + Vite 前端，Flask + SQLAlchemy 后端，REST JSON API。"),
            ("数据口径", "依据仓库 README、TECH_PLAN、技术选型与约定、API 设计文档及阶段总结整理。"),
            ("生成日期", "2026-06-27"),
        ],
    )
    add_callout(
        doc,
        "产品一句话",
        "PLEX Universe 通过学习任务、试炼挑战、知识星轨、成长档案和教师数据看板，把传统编程学习过程转化为可跟踪、可反馈、可激励的个性化学习闭环。",
    )

    add_heading(doc, "一、产品概述", 1)
    add_para(
        doc,
        "PLEX Universe（项目代号 A3）是一套面向编程学习场景的个性化学习系统。平台围绕“学生学习、教师指导、平台治理”三类核心场景设计，采用学生端、教师端和管理端三端分离的工作台结构。学生在平台中完成每日委托、班级试炼、知识星轨探索和成长档案复盘；教师通过领航总览、星域观测和 Explorer 档案观察班级状态、发现薄弱知识点并发布试炼；管理员负责系统策略、通知开关、试炼规则、AI 策略和数据安全配置。",
    )
    add_para(
        doc,
        "本产品的目标不是只做一个题库或静态课程展示页，而是把学习资源管理、AI 辅导、个性化推荐、学习效果评估和游戏化激励组合成一个连续系统。其核心价值在于让学生能看到“今天做什么、为什么做、做完有什么反馈”，让教师能看到“班级整体在哪里卡住、哪些学生需要跟进、试炼是否有效”。",
    )

    add_heading(doc, "二、建设背景与用户痛点", 1)
    add_matrix_table(
        doc,
        ["用户", "典型痛点", "产品回应"],
        [
            ("学生", "学习路径不清晰，练习反馈滞后，缺少持续学习动力。", "通过每日委托、星轨路径、试炼挑战、XP/等级/成就和成长档案形成即时反馈。"),
            ("教师", "难以及时掌握班级知识掌握情况，人工追踪学生进度成本高。", "通过教师聚合看板、热力图、排名、关注学生和 Explorer 档案提供班级诊断入口。"),
            ("管理员", "平台规则、通知、试炼和安全策略分散，难以统一管理。", "通过控制中枢集中维护规则配置、AI 教学策略、通知开关和数据安全选项。"),
            ("项目评审", "需要判断系统是否具备真实工程结构和可演示闭环。", "提供前后端分离、统一 API、角色权限、演示账号、种子数据和一键启动脚本。"),
        ],
        (0.9, 2.6, 3.0),
    )

    add_heading(doc, "三、产品定位与目标", 1)
    add_para(doc, "产品定位：面向课堂或训练营编程学习场景的个性化学习与教学管理平台。")
    add_para(doc, "总体目标：开发一套集学习资源管理、AI 智能辅导、个性化推荐、学习效果评估于一体的智能学习平台，提升学习积极性与学习效果。")
    add_para(doc, "阶段目标：当前项目已处于教学平台 MVP 迭代阶段，具备用户认证、三端路由、教师聚合看板、学生任务/试炼闭环、管理端配置页和基础测试覆盖。")
    add_heading(doc, "目标拆解", 2)
    for item in [
        "建立基础管理能力：用户管理、权限管理、班级分组管理和演示账号体系。",
        "建立学习激励能力：积分、等级、成就、排名、每日委托和试炼参与记录。",
        "建立教学观察能力：教师端按班级查看热力图、排名、关注学生、近期动态和试炼参与情况。",
        "建立个性化闭环能力：通过学生画像、资源生成、学习效果接口和错题/薄弱点记录支撑后续 AI 推荐。",
        "建立可交付工程能力：前后端分离、一键启动、统一响应、JWT 权限、种子数据和安全检查。 ",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "四、用户角色与使用场景", 1)
    add_matrix_table(
        doc,
        ["角色", "入口", "主要场景"],
        [
            ("学生", "/student 及子路由", "查看个人成长、领取今日委托、参与班级试炼、完成教师题目、探索知识星轨、查看成长档案。"),
            ("教师", "/teacher 及子路由", "查看班级总览、发布试炼、观察知识星域风险、追踪学生 Explorer 档案、管理班级学习节奏。"),
            ("管理员", "/admin", "维护系统设置、AI 策略、通知开关、试炼规则、平台配置与数据安全策略。"),
        ],
        (0.85, 1.45, 4.2),
    )
    add_heading(doc, "典型学习流程", 2)
    for step in [
        "学生登录后进入学生首页，查看等级、XP、班级排名、今日委托和试炼入口。",
        "教师在试炼中枢为班级发布试炼，系统按知识点生成或关联题目。",
        "学生在探索舱或今日委托中看到待做题目，提交答案后系统记录进度、积分和成就变化。",
        "教师在领航总览、星域观测和 Explorer 档案中查看班级整体与个体变化。",
        "管理员在控制中枢维护平台规则，保证演示和教学过程稳定可控。",
    ]:
        add_number(doc, step)

    add_heading(doc, "五、核心功能说明", 1)
    add_heading(doc, "5.1 学生端", 2)
    for item in [
        "学生首页：聚合 XP、等级、班级排名、今日委托、成就与成长摘要。",
        "探索舱：连接知识星域、班级试炼入口和补给站紧急任务，突出待修复碎片与学习行动。",
        "今日委托：展示每日任务进度、教师布置题目、完成反馈和奖励预览。",
        "学生试炼：查看、加入和完成教师发布的班级试炼，完成后进入积分和试炼记录闭环。",
        "探索档案：展示成长轨迹、技能分布、成就收藏、试炼与紧急任务记录。",
        "星轨路径：按语言基础、算法基础、动态规划、计算几何、图论、数据结构等学域导航知识点。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "5.2 教师端", 2)
    for item in [
        "领航总览：聚合班级指标、热力图、排名、今日委托完成率、关注学生和近期动态。",
        "星域观测：按知识领域观察班级风险和掌握情况，辅助教师定位薄弱点。",
        "Explorer 档案：查看学生画像、成长曲线、知识掌握、委托、试炼和成就记录。",
        "试炼中枢：创建、发布、编辑和查看班级试炼，支持 draft、scheduled、running、ended 等状态流转。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "5.3 管理端", 2)
    for item in [
        "控制中枢：维护班级系统设置、AI 教学策略、通知开关、试炼规则和数据安全设置。",
        "权限隔离：管理员可查看和维护全局配置，教师仅能维护本人班级相关配置。",
        "系统治理：为后续上线前的 License、CI、权限审计和环境变量治理保留入口。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "六、信息架构与页面路径", 1)
    add_matrix_table(
        doc,
        ["端口", "页面路径", "说明"],
        [
            ("学生端", "/student", "学生首页，总览个人学习状态。"),
            ("学生端", "/student/discovery", "探索舱，连接知识星域与试炼入口。"),
            ("学生端", "/student/daily", "今日委托，完成每日任务和教师题目。"),
            ("学生端", "/student/trials", "学生试炼列表与参与记录。"),
            ("学生端", "/student/archives", "探索档案，展示成长轨迹与成就。"),
            ("学生端", "/student/star-path", "星轨路径，六大学域知识点导航。"),
            ("教师端", "/teacher", "领航总览，班级指标与学习热力图。"),
            ("教师端", "/teacher/starfield", "星域观测，知识领域风险诊断。"),
            ("教师端", "/teacher/explorers", "Explorer 档案，学生画像与记录。"),
            ("教师端", "/teacher/trials", "试炼中枢，创建和发布班级试炼。"),
            ("管理端", "/admin", "控制中枢，维护平台策略和系统配置。"),
        ],
        (0.9, 1.85, 3.75),
    )

    add_heading(doc, "七、技术架构说明", 1)
    add_para(
        doc,
        "系统采用前后端分离架构。前端负责页面渲染、用户交互、状态管理和本地缓存；后端负责业务逻辑、数据操作、权限验证、AI 调用和数据统计；数据层以 MySQL 为生产主库，本地开发可使用 SQLite。开发环境中，Vite 将 `/api` 请求代理至 Flask 后端，统一 API 前缀为 `/api/v1`。",
    )
    add_matrix_table(
        doc,
        ["层级", "技术/组件", "职责"],
        [
            ("前端应用层", "Vue 3 + Vite + Naive UI + Pinia + Axios", "学生端、教师端、管理端页面，组件化交互和状态管理。"),
            ("API 层", "REST JSON + JWT", "统一响应结构、身份认证、角色权限和跨端数据契约。"),
            ("后端业务层", "Flask + SQLAlchemy + Marshmallow", "用户、班级、权限、试炼、激励、学习效果等业务处理。"),
            ("AI 能力层", "大模型/资源生成服务接口", "学生画像、资源生成、问答、课程安全检查和后续个性化推荐。"),
            ("数据存储层", "MySQL / SQLite，后续可接 Redis", "结构化数据、积分日志、排名缓存、试炼记录、系统配置。"),
            ("运维入口", "start.bat、测试脚本、种子脚本", "本地一键启动、环境检查、演示数据初始化和基础验证。"),
        ],
        (1.1, 2.15, 3.25),
    )
    add_callout(
        doc,
        "技术栈定案",
        "正式前端唯一目标栈为 Vue 3 + Vite + Naive UI + Pinia + Axios；后端为 Flask + SQLAlchemy；认证采用 JWT；接口采用 REST + JSON。React 原型仅作为视觉与交互参考，不作为主工程正式栈。",
    )

    add_heading(doc, "八、数据与接口设计", 1)
    add_para(
        doc,
        "后端接口采用统一响应结构：成功响应包含 `code: 0`、`message` 和 `data`；错误响应包含稳定错误码、错误信息和空数据。需要认证的请求通过 `Authorization: Bearer <access_token>` 携带 JWT。",
    )
    add_matrix_table(
        doc,
        ["领域", "主要数据表/实体", "说明"],
        [
            ("用户与权限", "users、roles、permissions、role_permissions", "支持学生、教师、管理员三类角色和 RBAC 权限控制。"),
            ("班级管理", "classes、用户班级关联", "支持教师班级、学生归属、班级统计和权限隔离。"),
            ("激励系统", "achievements、user_achievements、points_log、ranking_cache", "支持积分、等级、成就、连续学习和班级/周排名。"),
            ("试炼系统", "trials、trial_participations、trial_questions、trial_question_progress", "支持教师发布试炼、学生参与、题目下发和作答进度。"),
            ("紧急任务", "emergency_mission_sessions、emergency_mission_questions", "根据薄弱点生成补给站任务，记录作答与 XP 奖励。"),
            ("系统配置", "system_settings", "支持全局或班级级别规则、AI 策略、通知和安全设置。"),
        ],
        (1.0, 2.7, 2.8),
    )
    add_matrix_table(
        doc,
        ["接口类别", "代表接口", "用途"],
        [
            ("认证", "POST /api/v1/auth/login", "用户登录并获取 JWT。"),
            ("用户", "GET /api/v1/users/me", "获取当前用户资料、等级、积分、排名和激励信息。"),
            ("教师聚合", "GET /api/v1/teacher/overview", "教师端三类页面共享的班级总览事实源。"),
            ("试炼", "GET/POST/PATCH /api/v1/teacher/trials", "教师创建、发布、结束和查看试炼。"),
            ("学生试炼", "GET /api/v1/student/trials，POST /join，POST /complete", "学生查看、加入和完成班级试炼。"),
            ("学习路径", "GET /api/v1/student/learning-path", "返回六大学域进度与知识点路径。"),
            ("学习效果", "GET /api/v1/student/learning-effect", "基于资源介入前后作答证据评估学习效果。"),
            ("健康检查", "GET /api/v1/health，GET /api/v1/system/ai-health", "用于本地启动和系统状态验证。"),
        ],
        (1.0, 2.55, 2.95),
    )

    add_heading(doc, "九、界面与交互设计", 1)
    add_para(
        doc,
        "产品视觉方向采用“高活力、游戏化、沉浸式学习体验”。学生端强调成长、反馈和探索感；教师端强调班级洞察、风险识别和行动建议；管理端强调稳定、权威和可控。三端保持统一的 PLEX 宇宙风格，但通过不同主题色区分使用场景。",
    )
    add_matrix_table(
        doc,
        ["端口", "主题色", "视觉重点"],
        [
            ("学生端", "活力绿、青色、紫色", "成长、进度、成就、任务反馈和探索感。"),
            ("教师端", "暖橙、红色、黄色", "指导、预警、激励和班级洞察。"),
            ("管理端", "深紫、靛蓝、粉色", "配置、权限、治理和系统控制。"),
        ],
        (1.0, 1.9, 3.6),
    )
    add_para(
        doc,
        "游戏化机制包括等级 Lv1-Lv5、XP 积分、连续学习天数、成就勋章、班级排名、周排名、今日委托、试炼奖励和补给站任务。界面通过进度条、热力图、雷达图、成长曲线、任务卡片和成就反馈降低学习过程的抽象感。",
    )

    add_heading(doc, "十、运行与演示说明", 1)
    add_para(doc, "Windows 本地演示优先使用仓库根目录的 `start.bat`。双击脚本后会检查 Python、Node.js 和 npm，按需安装依赖，非破坏性初始化或升级数据库，分别启动 Flask 后端和 Vue 前端，并打开前端地址。")
    add_key_value_table(
        doc,
        [
            ("前端地址", "http://localhost:5173"),
            ("后端地址", "http://127.0.0.1:5000"),
            ("环境检查", "start.bat --check"),
            ("启动但不打开浏览器", "start.bat --no-browser"),
            ("前端构建", "cd frontend && npm run build"),
            ("后端测试", "cd backend && python -m unittest discover -s tests"),
        ],
    )
    add_matrix_table(
        doc,
        ["角色", "用户名", "密码", "入口"],
        [
            ("管理员", "admin", "admin123", "/admin"),
            ("教师", "teacher001", "teacher123", "/teacher"),
            ("学生", "student001", "student123", "/student"),
            ("学生演示数据", "explorer01 - explorer10", "student123", "/student"),
        ],
        (1.1, 1.8, 1.4, 2.2),
    )

    add_heading(doc, "十一、产品亮点", 1)
    for item in [
        "三端角色清晰：学生、教师、管理员各自拥有独立工作台，避免功能混杂。",
        "学习闭环完整：任务、试炼、作答、积分、成就、排名、档案和教师反馈形成连续链路。",
        "数据驱动教学：教师端通过热力图、排名、关注学生、星域观测和档案页辅助教学决策。",
        "工程结构真实：前后端分离、统一 API、JWT、RBAC、数据库模型、种子脚本和一键启动均已落地。",
        "个性化扩展明确：学生画像、资源生成、学习效果、错题本和知识库安全接口为后续 AI 能力预留基础。",
        "边界表达诚实：学习效果接口在证据不足时返回 `insufficient_evidence`，避免展示虚假的提升指标。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "十二、当前边界与后续规划", 1)
    add_para(
        doc,
        "当前项目已具备教学平台 MVP 的主体闭环，但仍有若干边界需要在正式生产部署前补齐。真实讯飞调用、在线 GitHub Actions/MySQL 8 结果仍属于外部条件阻塞；TTS 尚未实现；生产部署前还建议补齐数据库迁移、权限审计、License、CI 和更完整的环境变量示例。",
    )
    add_matrix_table(
        doc,
        ["方向", "当前状态", "后续建议"],
        [
            ("AI 资源生成", "已预留资源生成、画像、课程安全和学习效果接口。", "接入真实模型服务，补充教师审核、引用和失败恢复策略。"),
            ("试炼系统", "教师发布、学生参与、完成记录和奖励闭环已具备。", "完善真实关卡流程、完成结算、奖励发放和战绩接口。"),
            ("学习评估", "已要求前后测各至少 3 条有效作答，证据不足时不生成提升。", "扩大题目覆盖、错题本联动和班级维度分析。"),
            ("班级治理", "教师聚合看板、热力图、排名和关注学生已可演示。", "补齐数据导出、备份接口和更细权限审计。"),
            ("上线交付", "本地一键启动和演示账号可用。", "完善生产环境配置、CI 验证、MySQL 8 实测、License 和部署手册。"),
        ],
        (1.0, 2.7, 2.8),
    )

    add_heading(doc, "十三、结论", 1)
    add_para(
        doc,
        "PLEX Universe / A3 个性化学习系统已经形成较完整的网站产品形态：它不只包含页面展示，还包含三端角色、业务数据、权限控制、激励机制、教学看板、试炼闭环和本地运行路径。作为课程或项目报告中的网站产品说明，本系统可以被概括为“以编程学习为场景、以数据反馈为核心、以游戏化机制提升参与度、以教师诊断和 AI 个性化为后续增长点”的学习平台。",
    )
    add_para(
        doc,
        "从产品价值看，系统解决的是学习过程不可见、学习动力不足、教师诊断成本高和教学平台缺少连续反馈的问题；从工程实现看，系统采用明确的 Vue + Flask 前后端分离技术栈，具备可运行、可演示、可扩展的项目基础。后续只需围绕真实 AI 能力、生产部署验证和评审交付材料继续补齐，即可从 MVP 进一步推进到正式交付版本。",
    )

    add_heading(doc, "附录：文档依据", 1)
    for item in [
        "README.md：项目总体说明、三端功能、运行方式和当前状态。",
        "技术选型与约定.md：正式技术栈、Vue 定案、后端与认证约定。",
        "TECH_PLAN.md：技术方案总纲、阶段目标、核心模块和架构说明。",
        "backend_api_design.md：API 总体设计、数据表、核心接口、学习效果与安全接口。",
        "docs/2026-05-27-student-exploration-summary.md：学生探索链路、试炼题目、紧急任务和星轨路径沉淀。",
        "frontend_design_v2.md：前端 UI/UX、游戏化机制与三端视觉方向。",
    ]:
        add_bullet(doc, item)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(path.resolve())
